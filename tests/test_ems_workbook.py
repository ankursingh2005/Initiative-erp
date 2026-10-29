"""Excel formatting checks use synthetic report data only."""
import unittest
from io import BytesIO
from openpyxl import load_workbook
from ems_workbook import build_salary_workbook, IDS_HEADERS


class SalaryWorkbookTests(unittest.TestCase):
    def test_two_ordered_sheets_preserve_money_accounts_and_literal_names(self):
        ids = ['EMP001', '=Literal employee', 'Branch', 'Regular', 2, 30, 28,
               *([100.25] * 14), '000123456789', 'SBIN0001234']
        self.assertEqual(len(ids), len(IDS_HEADERS))
        phase2 = ['EMP001', '=Literal employee', 'Branch', 'Regular', 4, 2, 2,
                  100.25, 6.68, 500, 506.68]
        data = build_salary_workbook('2026-09', [ids], [phase2], True)
        self.assertTrue(data.startswith(b'PK'))
        workbook = load_workbook(BytesIO(data))
        self.assertEqual(workbook.sheetnames, ['IDS', 'W-O & Incentive'])
        first, second = workbook.worksheets
        self.assertEqual(first['V5'].value, '000123456789')
        self.assertEqual(first['V5'].data_type, 's')
        self.assertEqual(first['B5'].data_type, 's')
        self.assertEqual(first['H5'].value, 100.25)
        self.assertEqual(first['U6'].value, first['U5'].value)
        self.assertEqual(second['K5'].value, 506.68)
        self.assertEqual(second['K6'].value, 506.68)
        self.assertIn('Provisional', first['A3'].value)
        self.assertEqual(first.freeze_panes, 'E5')
        self.assertEqual(first.auto_filter.ref, 'A4:W5')


if __name__ == '__main__':
    unittest.main()
