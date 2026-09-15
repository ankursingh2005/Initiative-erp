"""Verify report logic without importing the application's database/server dependencies."""
import ast
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from types import SimpleNamespace
from typing import Optional
from io import BytesIO
import unittest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

source = Path('main.py').read_text(encoding='utf-8')
tree = ast.parse(source)
names = {'_incentive_header', '_incentive_outlet_short_name', 'build_non_sales_report'}
nodes = [node for node in tree.body if
         isinstance(node, ast.FunctionDef) and node.name in names or
         isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and
         t.id == 'NON_SALES_ALLOCATIONS' for t in node.targets)]
exec(compile(ast.Module(body=nodes, type_ignores=[]), 'main.py', 'exec'))
import re

# Run the same screenshot, missing-data, zero and negative rounding regression.
tests = ast.parse(Path('tests/test_ids_fund.py').read_text(encoding='utf-8'))
case = next(n for n in tests.body if isinstance(n, ast.ClassDef))
case.body = [n for n in case.body if isinstance(n, ast.FunctionDef) and
             n.name == 'test_non_sales_screenshot_and_pending']
main = SimpleNamespace(build_non_sales_report=build_non_sales_report)
exec(compile(ast.Module(body=[case], type_ignores=[]), 'test_ids_fund.py', 'exec'))
result = unittest.TextTestRunner().run(unittest.defaultTestLoader.loadTestsFromTestCase(IdsFundTests))
assert result.wasSuccessful()

book = Workbook()
summary = book.active
summary.title = 'IDS Fund Outlet Summary'
fund_report = {'summary': [dict(outlet=code, ids_fund=value) for code, value in
    zip(['ALM', 'ASH', 'HZT', 'GNG', 'VKN'], [11479, 837, 7214, 2166, 1675])]}
for index, row in enumerate(fund_report['summary'], 3):
    summary.cell(index, 2, row['ids_fund'])
navy, white, pale, thin = '17365D', 'FFFFFF', 'D9EAF7', Side(style='thin')
start = source.index('    staff = book.create_sheet("Non-sales Staff Incentive")')
end = source.index('    output = BytesIO()', start)
import textwrap
sheet, category_sheet, group_sheet, exact_sheet, fund_sheet, outlet_fund = [
    book.create_sheet(name) for name in ['Main', 'Category', 'Group', 'Exact', 'Fund', 'Outlet']]
exact_sheet.append(['Title'])
exact_sheet.append(['Note'])
exact_sheet.append(['Outlet', 'Category Group', 'Total Incentive', 'Applied Rate', 'Exact Incentive'])
exact_sheet.append(['ALM', 'HA + HE', 100, .85, 85])
exec(textwrap.dedent(source[start:end]))
for column in 'CDE':
    assert exact_sheet[column+'3'].alignment.horizontal == exact_sheet[column+'4'].alignment.horizontal == 'right'
assert staff['C11'].value == '=IF(ISNUMBER(B2),ROUND(B2*C2,0),"Pending")'
assert staff['D16'].value == '=IF(ISNUMBER(B7),ROUND(B7*D7,2),"Pending")'
assert staff['D18'].value == '=D16'
assert staff['B2'].value == "='IDS Fund Outlet Summary'!B3"
assert staff['B7'].value == '=IF(COUNT(B2:B6)=5,SUM(B2:B6),"Pending")'
assert staff['G18'].value == '=G17'
buffer = BytesIO()
book.save(buffer)
buffer.seek(0)
reloaded = load_workbook(buffer)
assert reloaded['Non-sales Staff Incentive']['D7'].value == .16
assert reloaded['Non-sales Staff Incentive']['C2'].number_format == '0%'
print('Screenshot calculations, rounding, pending values and Excel generation passed.')

# Exercise the selected export handler without server/database startup.
handler = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == 'export_selected_incentives')
handler.decorator_list = []
handler.args.defaults = []
for argument in handler.args.args:
    argument.annotation = None
exec(compile(ast.Module(body=[handler], type_ignores=[]), 'main.py', 'exec'))
from datetime import date
india_today = date.today
Response = lambda content, **kwargs: SimpleNamespace(content=content, **kwargs)
calculate_incentive_report = lambda *args: {
    'exact_summary': [dict(outlet='ALM', group='HA + HE', total_incentive=100, applied_rate=85, exact_incentive=85)],
    'ids_fund_report': {'non_sales_report': build_non_sales_report(fund_report['summary'])}}
for selection, sheets in [('sales', ['Sales Incentive']), ('non-sales', ['Non-sales Incentive']),
                          ('both', ['Sales Incentive', 'Non-sales Incentive'])]:
    for format in ['xlsx', 'pdf']:
        response = export_selected_incentives(None, 7, 2.5, selection, format, None)
        if format == 'xlsx':
            exported = load_workbook(BytesIO(response.content))
            assert exported.sheetnames == sheets
            if 'Non-sales Incentive' in sheets:
                assert exported['Non-sales Incentive']['D9'].value == 3739.36
                for column in 'BCDEFG':
                    assert exported['Non-sales Incentive'][column+'2'].alignment.horizontal == 'right'
                    assert exported['Non-sales Incentive'][column+'3'].alignment.horizontal == 'right'
            if 'Sales Incentive' in sheets:
                assert exported['Sales Incentive']['C3'].value == 85
                assert exported['Sales Incentive'].max_column == 3
                assert [c.value for c in exported['Sales Incentive'][2]] == ['Outlet', 'Category Group', 'Exact Incentive']
        else:
            assert response.content.startswith(b'%PDF')
print('Individual and combined Excel/PDF exports passed.')
handler = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == 'export_exact_incentive_report')
handler.decorator_list = []
handler.args.defaults = []
for argument in handler.args.args:
    argument.annotation = None
exec(compile(ast.Module(body=[handler], type_ignores=[]), 'main.py', 'exec'))
response = export_exact_incentive_report(None, 7, 2.5, 'xlsx', None)
exported = load_workbook(BytesIO(response.content))
assert exported.active.max_column == 3
assert exported.active['C3'].value == 85
print('Exact report download also contains only the three sales columns.')
