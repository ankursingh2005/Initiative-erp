"""Isolated authorization regressions for manager outlet attendance."""
from datetime import date, datetime
from io import BytesIO
import unittest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from openpyxl import load_workbook
import test_analytics as isolated
import auth
import models
from attendance_exports import router

main = isolated.main
tearDownModule = isolated.tearDownModule


class ManagerAttendanceTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
        models.Base.metadata.create_all(self.engine)
        self.db = sessionmaker(bind=self.engine)()
        self.stores = [models.Store(name='Outlet A'), models.Store(name='Outlet B')]
        self.db.add_all(self.stores)
        self.db.flush()
        self.users = []
        for index, (role, outlet) in enumerate([('CategoryManager', 0), ('Employee', 0), ('Employee', 1), ('CategoryManager', None), ('Admin', None)]):
            user = models.User(username='user'+str(index), email=str(index)+'@example.test', role=role,
                               password_hash='unused', status='Active', store_id=self.stores[outlet].id if outlet is not None else None)
            self.db.add(user)
            self.users.append(user)
        self.db.flush()
        self.records = []
        for employee, outlet in [(1, 0), (2, 1), (1, 1)]:
            record = models.AttendanceRecord(user_id=self.users[employee].id, store_id=self.stores[outlet].id,
                attendance_date=date(2026, 9, 23 if outlet == 0 else 22), checkin_at=datetime(2026, 9, 23 if outlet == 0 else 22, 9), checkin_selfie='saved-photo')
            self.db.add(record)
            self.records.append(record)
        self.db.commit()
        self.actor = self.users[0]
        app = FastAPI()
        app.include_router(router)
        app.get('/summary')(main.attendance_admin_summary)
        app.get('/profile', response_model=main.schemas.MyProfileOut)(main.get_my_profile)
        app.get('/history/{user_id}')(main.attendance_user_history)
        app.get('/selfies/{record_id}')(main.get_attendance_selfies)
        app.dependency_overrides[main.get_db] = lambda: self.db
        app.dependency_overrides[auth.get_current_user] = lambda: self.actor
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        self.db.close()
        self.engine.dispose()

    def test_dashboard_defaults_to_own_outlet_and_rejects_override(self):
        response = self.client.get('/summary?from_date=2026-09-23&to_date=2026-09-23')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual({row['user_id'] for row in response.json()['rows']}, {self.users[0].id, self.users[1].id})
        self.assertEqual(self.client.get('/summary?store_id='+str(self.stores[1].id)).status_code, 403)
        self.actor = self.users[3]
        self.assertEqual(self.client.get('/summary').status_code, 403)
        self.actor = self.users[1]
        self.assertEqual(self.client.get('/summary').status_code, 403)
        self.actor = self.users[4]
        self.assertEqual(self.client.get('/summary?store_id='+str(self.stores[1].id)).status_code, 200)

    def test_history_and_photos_are_scoped_and_profiles_remain_read_only(self):
        response = self.client.get('/history/'+str(self.users[1].id))
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json()['can_edit_profile'])
        self.assertEqual([row['id'] for row in response.json()['history']], [self.records[0].id])
        self.assertEqual(self.client.get('/history/'+str(self.users[2].id)).status_code, 403)
        self.assertEqual(self.client.get('/selfies/'+str(self.records[0].id)).status_code, 200)
        for record in self.records[1:]:
            self.assertEqual(self.client.get('/selfies/'+str(record.id)).status_code, 403)

    def test_exports_cannot_leak_other_outlet_users_or_records(self):
        for path in ['/api/attendance/admin-export?date=2026-09-23', '/api/attendance/monthly-export?month=2026-09']:
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else '')
            rows = list(load_workbook(BytesIO(response.content)).active.values)[1:]
            self.assertTrue(rows)
            self.assertTrue(all(row[0] in {self.users[0].id, self.users[1].id} for row in rows))
            self.assertEqual(self.client.get(path+'&store_id='+str(self.stores[1].id)).status_code, 403)
        self.assertEqual(self.client.get('/api/attendance/admin-user-export?user_id='+str(self.users[2].id)).status_code, 403)
        response = self.client.get('/api/attendance/admin-user-export?user_id='+str(self.users[1].id))
        self.assertEqual(response.status_code, 200)
        rows = list(load_workbook(BytesIO(response.content)).active.values)[1:]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][3], 'Outlet A')

    def test_manager_only_sees_ids_and_brand_categories(self):
        brand = models.User(username='Promoter', email='promoter@example.test', role='BrandPartner', store_id=self.stores[0].id, status='Active', password_hash='unused')
        ac = models.User(username='AC', email='ac@example.test', role='AC Technician A', store_id=self.stores[0].id, status='Active', password_hash='unused')
        assigned = models.User(username='Assigned AC', email='akhtarnoor3112@gmail.com', role='Service Manager A', store_id=self.stores[0].id, status='Active', password_hash='unused')
        self.db.add_all([brand, ac, assigned])
        self.db.commit()
        response = self.client.get('/summary')
        self.assertEqual(response.status_code, 200)
        ids = {row['user_id'] for row in response.json()['rows']}
        self.assertIn(brand.id, ids)
        self.assertNotIn(ac.id, ids)
        self.assertNotIn(assigned.id, ids)
        for user in (ac, assigned):
            self.assertEqual(self.client.get('/history/'+str(user.id)).status_code, 403)
            self.assertEqual(self.client.get('/api/attendance/admin-user-export?user_id='+str(user.id)).status_code, 403)
        for category in ('ac_retails', 'ac_projects'):
            self.assertEqual(self.client.get('/summary?emp_category='+category).status_code, 403)
            self.assertEqual(self.client.get('/api/attendance/admin-export?date=2026-09-23&emp_category='+category).status_code, 403)
            self.assertEqual(self.client.get('/api/attendance/monthly-export?month=2026-09&emp_category='+category).status_code, 403)

    def test_summary_uses_latest_saved_profile_name_and_email(self):
        employee = self.users[1]
        self.db.add(models.IdentityCard(user_id=employee.id, employee_id='IDS-TEST-26001', employee_name='MAHESH PRATAP SINGH', designation='Employee', mobile='9876543210'))
        employee.full_name = 'MAHESH PRATAP SINGH'
        employee.email = 'updated@example.test'
        self.db.commit()
        for actor in (self.users[0], self.users[4]):
            self.actor = actor
            response = self.client.get('/summary')
            self.assertEqual(response.status_code, 200)
            row = next(item for item in response.json()['rows'] if item['user_id'] == employee.id)
            self.assertEqual(row['display_name'], 'MAHESH PRATAP SINGH')
            self.assertEqual(row['email'], 'updated@example.test')
            self.assertEqual(row['username'], employee.username)


    def test_noor_service_manager_is_limited_to_ac_projects(self):
        noor = models.User(username='Noor Akhtar', email='akhtarnoor3112@gmail.com', role='Service Manager A', status='Active', password_hash='unused')
        project = models.User(username='Project tech', email='project@example.test', role='AC Technician A', status='Active', password_hash='unused', store_id=self.stores[1].id)
        retail = models.User(username='Retail tech', email='retail@example.test', role='AC Technician B', status='Active', password_hash='unused')
        other = models.User(username='Other service manager', email='other@example.test', role='Service Manager B', status='Active', password_hash='unused')
        self.db.add_all([noor, project, retail, other])
        self.db.flush()
        record = models.AttendanceRecord(user_id=project.id, store_id=self.stores[1].id, attendance_date=date(2026, 9, 23), checkin_at=datetime(2026, 9, 23, 9), checkin_selfie='project-photo')
        self.db.add(record)
        self.db.commit()
        self.actor = noor
        response = self.client.get('/summary?from_date=2026-09-23&to_date=2026-09-23')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual({row['user_id'] for row in response.json()['rows']}, {noor.id, project.id})
        self.assertEqual(self.client.get('/history/'+str(project.id)).status_code, 200)
        self.assertFalse(self.client.get('/history/'+str(project.id)).json()['can_edit_profile'])
        self.assertEqual(self.client.get('/selfies/'+str(record.id)).status_code, 200)
        for user in (retail, other, self.users[1]):
            self.assertEqual(self.client.get('/history/'+str(user.id)).status_code, 403)
            self.assertEqual(self.client.get('/api/attendance/admin-user-export?user_id='+str(user.id)).status_code, 403)
        self.assertEqual(self.client.get('/selfies/'+str(self.records[0].id)).status_code, 403)
        for path in ('/summary', '/api/attendance/admin-export?date=2026-09-23', '/api/attendance/monthly-export?month=2026-09'):
            for category in ('ids_emp', 'brand_pro', 'ac_retails'):
                self.assertEqual(self.client.get(path+('&' if '?' in path else '?')+'emp_category='+category).status_code, 403)
        for path in ('/api/attendance/admin-export?date=2026-09-23', '/api/attendance/monthly-export?month=2026-09'):
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200)
            rows = list(load_workbook(BytesIO(response.content)).active.values)[1:]
            self.assertTrue(rows)
            self.assertTrue(all(row[0] in {noor.id, project.id} for row in rows))
        self.actor = other
        self.assertEqual(self.client.get('/summary').status_code, 200)
        self.assertEqual(self.client.get('/history/'+str(project.id)).status_code, 403)
        self.assertEqual(self.client.get('/api/attendance/admin-export?date=2026-09-23').status_code, 200)
        self.actor = noor
        noor.role = 'Employee'
        self.db.commit()
        self.assertEqual(self.client.get('/summary').status_code, 403)


    def test_profile_exposes_service_dashboard_permission_through_response_schema(self):
        for role, email, allowed in (
                ('Service Manager A', 'akhtarnoor3112@gmail.com', True),
                ('Service Manager A', 'other@example.test', True),
                ('Employee', 'akhtarnoor3112@gmail.com', False)):
            with self.subTest(role=role, email=email):
                self.actor.role = role
                self.actor.email = email
                self.db.commit()
                response = self.client.get('/profile')
                self.assertEqual(response.status_code, 200, response.text)
                self.assertIs(response.json()['ac_project_dashboard'], allowed)


    def test_leadership_can_view_all_outlets_details_and_exports(self):
        for role in ('CEO', 'Director', 'AccountsManager'):
            with self.subTest(role=role):
                self.actor = self.users[0]
                self.actor.role = role
                self.db.commit()
                response = self.client.get('/summary')
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual({row['user_id'] for row in response.json()['rows']}, {user.id for user in self.users})
                response = self.client.get('/summary?store_id='+str(self.stores[1].id))
                self.assertEqual(response.status_code, 200)
                self.assertEqual({row['user_id'] for row in response.json()['rows']}, {self.users[2].id})
                response = self.client.get('/history/'+str(self.users[2].id))
                self.assertEqual(response.status_code, 200)
                self.assertFalse(response.json()['can_edit_profile'])
                self.assertEqual(self.client.get('/selfies/'+str(self.records[1].id)).status_code, 200)
                for path in ('/api/attendance/admin-export?date=2026-09-22', '/api/attendance/monthly-export?month=2026-09', '/api/attendance/admin-user-export?user_id='+str(self.users[2].id)):
                    response = self.client.get(path)
                    self.assertEqual(response.status_code, 200, path)
                    rows = list(load_workbook(BytesIO(response.content)).active.values)[1:]
                    self.assertIn(self.users[2].id, {row[0] for row in rows})
