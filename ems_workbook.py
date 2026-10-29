"""Format EMS report snapshots as Excel; never read or write ERP records."""
from io import BytesIO

IDS_HEADERS = ['Employee ID', 'Employee Name', 'Branch', 'Category', 'Absent',
    'No. of Days', 'Working Days', 'Expected Basic', 'Actual Basic', 'Other Allowances',
    'Actual VDA', 'Actual Pay (EPF)', 'Actual OA', 'Total Earning', 'Bonus Payable',
    'Actual Gross Earning', 'EPF Employee Contribution', 'ESIC Employee Contribution',
    'Staff Advance & Deduction', 'Total Deductions', 'Net Amount', 'Account No.', 'IFSC Code']
WEEKOFF_HEADERS = ['Employee ID', 'Employee Name', 'Branch', 'Category', 'Total W/O',
    'W/O Availed', 'W/O Credit', 'Gross Salary', 'W/O Credit Pay', 'Incentive', 'Amount Payable']


def build_salary_workbook(month, ids_rows, weekoff_rows, provisional=False):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, title, headers, rows, money_columns in (
        ('IDS', 'IDS - Phase 1 Salary', IDS_HEADERS, ids_rows, range(8, 22)),
        ('W-O & Incentive', 'W/O & Incentive - Phase 2', WEEKOFF_HEADERS, weekoff_rows, range(8, 12)),
    ):
        sheet = workbook.create_sheet(name)
        sheet.append([f'Initiative Data Systems Pvt. Ltd. | {title}'])
        sheet.append([f'Salary month: {month}'])
        sheet.append(['Provisional: open month, incomplete attendance or unsaved inputs'
                      if provisional else 'Current EMS report snapshot'])
        sheet.append(headers)
        for row_number in (1, 2, 3):
            sheet.merge_cells(start_row=row_number, start_column=1,
                              end_row=row_number, end_column=len(headers))
        for values in rows:
            sheet.append(values)
        last_data = sheet.max_row
        totals = ['Total'] + [None] * (len(headers) - 1)
        for column in money_columns:
            totals[column - 1] = round(sum(values[column - 1] or 0 for values in rows), 2)
        sheet.append(totals)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    # Names, IDs and account numbers must remain literal text.
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='center')
                if cell.row >= 5 and cell.column in money_columns:
                    cell.number_format = '#,##0.00;[Red](#,##0.00);"-"'
                if cell.row in (1, 4, sheet.max_row):
                    cell.font = Font(bold=True, color='FFFFFF' if cell.row == 4 else '153349')
                    cell.fill = PatternFill('solid', fgColor='19485A' if cell.row == 4 else 'EAF3F5')
                if cell.row == 4:
                    cell.alignment = Alignment(wrap_text=True, vertical='center')
        for column, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(column)].width = min(29, max(16, len(header) + 2))
        sheet.column_dimensions['B'].width = 28
        sheet.row_dimensions[4].height = 42
        sheet.freeze_panes = 'E5'
        sheet.auto_filter.ref = f'A4:{get_column_letter(len(headers))}{last_data}'
        sheet.sheet_view.showGridLines = False
        sheet.print_title_rows = '1:4'
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
