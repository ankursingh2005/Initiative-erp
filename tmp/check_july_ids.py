import sys
sys.path.insert(0, '.')
sys.path.insert(0, 'tests')
import test_ids_fund as test
from pathlib import Path
from io import BytesIO
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook
import json
content = Path('C:/Users/Ankur/Desktop/JULY.xlsx').read_bytes()
parsed = test.main.parse_incentive_workbook(content, 'JULY.xlsx')
report = test.main.build_ids_fund_report(parsed)
book = load_workbook(BytesIO(content), data_only=True)
ws = book.active
# Independent computation directly from the reference cell positions.
expected = {}
for name, first_row in [('ALM',7),('ASH',17),('HZT',27),('GNG',37),('VKN',47)]:
    incentive = Decimal(0)
    fund = Decimal(0)
    sales_total = Decimal(0)
    for offset in range(7):
        category = ws.cell(first_row+offset,2).value
        sales = Decimal(str(ws.cell(first_row+offset,3).value))
        fraction = Decimal('.00125') if category=='MOB' else Decimal('.0025')
        if category in ('ACC','ACC COM'):
            fraction = Decimal('.01') if name!='HZT' else Decimal('.005') if category=='ACC' else Decimal('.0025')
        line_incentive = (sales*fraction).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
        incentive += line_incentive
        share = {'MOB':Decimal('.30'),'COM':Decimal('.20'),'DC':Decimal('.25'),'HA':Decimal('.20'),'HE':Decimal('.25'),'ACC COM':Decimal('.25'),'ACC':Decimal('.25')}[category]
        fund += (line_incentive*share).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
        sales_total += sales
    assert sales_total == Decimal(str(ws.cell(first_row+7,3).value))
    expected[name] = dict(total_sales=float(sales_total),total_incentive=float(incentive),ids_fund=float(fund))
for row in report['summary']:
    assert {k:row[k] for k in expected[row['outlet']]} == expected[row['outlet']], row
assert report['totals']['total_sales']==ws['C55'].value==42544283
assert report['pending_rows']==0
assert len(report['rows'])==35
print(json.dumps({'summary':report['summary'],'totals':report['totals'],'category_rows':len(parsed),'independent_reconciliation':'passed'},indent=2))
book.close()
test.tearDownModule()
