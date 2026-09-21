"""Attendance downloads built from saved user IDs and attendance records."""
import calendar
from datetime import date, timedelta
from io import BytesIO
from xml.sax.saxutils import escape
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy.orm import Session
import auth
import models
from database import get_db

router = APIRouter(prefix='/api/attendance')
HEADERS = ['User ID', 'Employee', 'Date', 'Outlet', 'Status', 'Punch in', 'Punch out', 'Hours']


def render_export(rows, format, filename):
    if format not in {'xlsx', 'pdf'}:
        raise HTTPException(400, 'Choose Excel or PDF')
    output = BytesIO()
    if format == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Attendance'
        sheet.append(HEADERS)
        for row in rows:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Names must remain text, never spreadsheet formulas.
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='155EEF')
        for index, width in enumerate([12, 30, 15, 25, 14, 15, 15, 12], 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        workbook.save(output)
        media = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        styles = getSampleStyleSheet()
        body = styles['BodyText']
        body.fontSize = 8
        body.leading = 10
        data = [[Paragraph(escape(str(value)), body) for value in row] for row in [HEADERS] + rows]
        table = Table(data, colWidths=[42, 155, 70, 125, 65, 65, 65, 50], repeatRows=1)
        table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EAF1FF')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('GRID', (0, 0), (-1, -1), .3, colors.lightgrey),
            ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))
        SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=30, rightMargin=30,
            topMargin=25, bottomMargin=25).build([Paragraph('Attendance report', styles['Title']), Spacer(1, 12), table])
        media = 'application/pdf'
    return Response(output.getvalue(), media_type=media,
                    headers={'Content-Disposition': f'attachment; filename="{filename}.{format}"'})


def export_rows(db, start=None, end=None, user_id=None, store_id=None, weekoff_day=None, status=None, emp_category=None):
    users = db.query(models.User)
    if user_id is not None:
        users = users.filter(models.User.id == user_id)
    else:
        users = users.filter(models.User.status == 'Active')
    if weekoff_day:
        users = users.filter(models.User.weekoff_day == weekoff_day)
    if emp_category == 'brand_pro':
        users = users.filter(models.User.role == 'BrandPartner')
    elif emp_category == 'ids_emp':
        users = users.filter(~models.User.role.in_(['BrandPartner', 'ACTechnicianA', 'ACTechnicianB']))
    elif emp_category:
        raise HTTPException(400, 'This employee category is not configured for export. Select all categories.')
    users = users.order_by(models.User.username, models.User.id).all()
    ids = [user.id for user in users]
    records = db.query(models.AttendanceRecord).filter(models.AttendanceRecord.user_id.in_(ids))
    leaves = db.query(models.AttendanceLeave).filter(models.AttendanceLeave.user_id.in_(ids))
    if start:
        records = records.filter(models.AttendanceRecord.attendance_date >= start, models.AttendanceRecord.attendance_date <= end)
        leaves = leaves.filter(models.AttendanceLeave.leave_date >= start, models.AttendanceLeave.leave_date <= end)
    by_day = {(r.user_id, r.attendance_date): r for r in records.order_by(models.AttendanceRecord.id).all()}
    leave_days = {(r.user_id, r.leave_date) for r in leaves.all()}
    stores = {s.id: s.name for s in db.query(models.Store).all()}
    rows = []
    for user in users:
        days = ([start + timedelta(days=i) for i in range((end-start).days+1)] if start else
                sorted({day for uid, day in set(by_day) | leave_days if uid == user.id}))
        for day in days:
            record = by_day.get((user.id, day))
            outlet_id = record.store_id if record else user.store_id
            if store_id is not None and outlet_id != store_id:
                continue
            state = ('Present' if record and record.checkin_at else 'Leave' if (user.id, day) in leave_days
                     else 'Week Off' if day.strftime('%A') == user.weekoff_day else 'Absent')
            if status and state != status:
                continue
            checkin, checkout = (record.checkin_at, record.checkout_at) if record else (None, None)
            hours = round((checkout-checkin).total_seconds()/3600, 2) if checkin and checkout else ''
            rows.append([user.id, user.username, day.isoformat(), stores.get(outlet_id, ''), state,
                         checkin.strftime('%H:%M') if checkin else '', checkout.strftime('%H:%M') if checkout else '', hours])
    return rows


@router.get('/admin-export')
def daily_export(date: date, format: str = 'xlsx', store_id: int | None = None,
                 weekoff_day: str | None = None, status: str | None = None, emp_category: str | None = None,
                 db: Session = Depends(get_db), actor=Depends(auth.require_roles('Admin'))):
    rows = export_rows(db, date, date, store_id=store_id, weekoff_day=weekoff_day, status=status, emp_category=emp_category)
    return render_export(rows, format, f'attendance-{date}')


@router.get('/monthly-export')
def monthly_export(month: str, store_id: int | None = None, weekoff_day: str | None = None,
                   status: str | None = None, emp_category: str | None = None,
                   db: Session = Depends(get_db), actor=Depends(auth.require_roles('Admin'))):
    try:
        start = date.fromisoformat(month + '-01')
    except ValueError:
        raise HTTPException(400, 'Month must be YYYY-MM')
    end = start.replace(day=calendar.monthrange(start.year, start.month)[1])
    return render_export(export_rows(db, start, end, store_id=store_id, weekoff_day=weekoff_day,
                         status=status, emp_category=emp_category), 'xlsx', f'attendance-{month}')


@router.get('/admin-user-export')
def user_export(user_id: int, format: str = 'xlsx', db: Session = Depends(get_db),
                actor=Depends(auth.require_roles('Admin'))):
    if not db.query(models.User).filter(models.User.id == user_id).first():
        raise HTTPException(404, 'User not found')
    return render_export(export_rows(db, user_id=user_id), format, f'attendance-user-{user_id}')
